from pptx import Presentation
from pptx.util import Inches, Pt

# Create presentation
prs = Presentation()

# Title Slide
slide_layout = prs.slide_layouts[0] 
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "NumPy: Random Generation & Slicing"
subtitle.text = "From Basics to 3D Data Manipulation"

# Slide 1: Why NumPy?
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "1. The Basics: Why NumPy?"
tf = body.text_frame
tf.text = "Python lists are great, but..."
p = tf.add_paragraph()
p.text = "NumPy is written in C, making it incredibly fast for large data."
p.level = 1
p = tf.add_paragraph()
p.text = "Uses much less memory than standard lists."
p.level = 1
p = tf.add_paragraph()
p.text = "import numpy as np"
p.level = 0
p.font.bold = True

# Slide 2: Generating Random Data (1D)
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "2. Generating Random Data (1D)"
tf = body.text_frame
tf.text = "Simulating a list of 10 student scores:"
p = tf.add_paragraph()
p.text = "scores = np.random.randint(1, 50, 10)"
p.level = 0
p.font.bold = True
p = tf.add_paragraph()
p.text = "Syntax: np.random.randint(low, high, size)"
p.level = 1
p = tf.add_paragraph()
p.text = "Note: The upper bound (50) is exclusive!"
p.level = 1

# Slide 3: Generating Random Data (2D)
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "3. Generating Random Data (2D)"
tf = body.text_frame
tf.text = "Simulating a data table (10 students, 5 subjects):"
p = tf.add_paragraph()
p.text = "data_table = np.random.randint(2, 12, (10, 5))"
p.level = 0
p.font.bold = True
p = tf.add_paragraph()
p.text = "Size (10, 5) means 10 Rows and 5 Columns."
p.level = 1
p = tf.add_paragraph()
p.text = "Use data_table.shape to verify dimensions."
p.level = 1

# Slide 4: Basic Indexing
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "4. Basic Indexing"
tf = body.text_frame
tf.text = "1D Indexing:"
p = tf.add_paragraph()
p.text = "scores[0]    # First item"
p.level = 1
p = tf.add_paragraph()
p.text = "scores[-1]   # Last item (Negative indexing)"
p.level = 1
p = tf.add_paragraph()
p.text = "2D Indexing [row, column]:"
p.level = 0
p = tf.add_paragraph()
p.text = "data_table[0, 1]   # 1st row, 2nd column"
p.level = 1

# Slide 5: 1D Slicing
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "5. Slicing 1D Arrays"
tf = body.text_frame
tf.text = "The Golden Rule: [start : stop : step]"
p = tf.add_paragraph()
p.text = "scores[2:5]   # Index 2 up to 4"
p.level = 1
p = tf.add_paragraph()
p.text = "scores[7:]    # Index 7 to the end"
p.level = 1
p = tf.add_paragraph()
p.text = "scores[::-1]  # Magic Trick: Reverses the array!"
p.level = 1

# Slide 6: 2D Slicing
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "6. Slicing 2D Arrays (Matrices)"
tf = body.text_frame
tf.text = "Syntax: [row_slice, column_slice]"
p = tf.add_paragraph()
p.text = "data_table[:, :2]"
p.level = 1
p.font.bold = True
p = tf.add_paragraph()
p.text = "All rows, First 2 columns"
p.level = 2
p = tf.add_paragraph()
p.text = "data_table[3:6, 2:]"
p.level = 1
p.font.bold = True
p = tf.add_paragraph()
p.text = "Rows 3-5, Columns from index 2 to end"
p.level = 2
p = tf.add_paragraph()
p.text = "data_table[0::2, 2:]"
p.level = 1
p.font.bold = True
p = tf.add_paragraph()
p.text = "Every alternate row, Columns from index 2 to end"
p.level = 2

# Slide 7: 3D Slicing (Advanced)
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "7. 3D Slicing (The Mind Bender)"
tf = body.text_frame
tf.text = "Imagine 5 spreadsheets stacked together:"
p = tf.add_paragraph()
p.text = "cube_data = np.random.randint(2, 50, (5, 10, 5))"
p.level = 1
p.font.bold = True
p = tf.add_paragraph()
p.text = "Syntax: [depth, row, column]"
p.level = 0
p = tf.add_paragraph()
p.text = "cube_data[:, 9:, :]"
p.level = 1
p.font.bold = True
p = tf.add_paragraph()
p.text = "Extracts the last row from ALL 5 spreadsheets!"
p.level = 2

prs.save('c:/Users/iamab/Desktop/LearnLog/NumPy_Class_Presentation.pptx')
print("Presentation created successfully!")
